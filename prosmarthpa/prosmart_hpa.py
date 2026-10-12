#*********************************************************** ProSmart HPA controller (paper Section 4, Algorithms 3 and 4) ***********************************************************
#
# The released ProSmart package cannot run (wrong argument count in Analyse, pd.concat on scalars, ARM predicting
# without regressors, "productcatalogueservice" typo, strings stored in np.zeros arrays, circular imports, Main.py
# calling an undefined "redis" function). This file re-implements Algorithms 3 and 4 of the paper instead, and keeps
# everything else identical to our Smart HPA runs so the comparison is fair:
#   * Monitor: the unchanged Monitor() of each Smart HPA Microservice Manager in ..\smarthpa (same kubectl calls,
#     same CPU averaging, same retry patch from config.py)
#   * scenario / duration: SMARTHPA_SCENARIO and SMARTHPA_DURATION through config.py
#   * Knowledge Base: same file per microservice and same columns A-F, so make_run_report.py works unchanged
#
# Knowledge Base columns:
#   A Test Time (sec)  B CPU Usage Percentage  C Current Replicas
#   D Desired Replicas      reactive DR = ceil(CR x CPU% / threshold), the SAME demand definition as Smart HPA and the
#                           K8s HPA observer, so the Table 1 metrics compare like with like (the original ProSmart
#                           code stored its own prediction here, which would make the metrics depend on the model)
#   E Max. Replicas         capacity in effect this iteration (PmaxR when the ARM ran, otherwise SLA maxR)
#   F Scaling Action        decision executed
#   G Predicted Desired Replicas (PDR)   H Executed Replicas   I Predicted Time (sec)   J ARM Active
#
# Run from ..\smarthpa (run_experiment.ps1 does this) with D:\prosmart-venv\Scripts\python.exe.

import math
import multiprocessing
import os
import pickle
import subprocess
import sys
import time

SMARTHPA_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "smarthpa")
MODEL_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "models")
sys.path.insert(0, SMARTHPA_DIR)

import pandas as pd
from openpyxl import load_workbook

import config                                                   # scenario values + kubectl retry patch

# Smart HPA module name -> deployment / Knowledge Base name
MANAGERS = {"frontend": "frontend", "adservice": "adservice", "cartservice": "cartservice",
            "currencyservice": "currencyservice", "checkoutservice": "checkoutservice", "emailservice": "emailservice",
            "paymentservice": "paymentservice", "shippingservice": "shippingservice",
            "productcatalogservice": "productcatalogservice", "recommendationservice": "recommendationservice",
            "rediscart": "redis-cart"}

PREDICTION_WINDOW_SEC = 25          # best window in the paper (Section 5.3)
PREDICTION_STEP_SEC = 5             # models are trained on a 5 s grid
PREDICTION_START_SEC = 10           # as in the released code: predict only after 10 s of load


def load_model(service, kind):
    suffix = "Demand_Model" if kind == "demand" else "capacity_model"
    with open(os.path.join(MODEL_DIR, f"{service}_{suffix}.pkl"), "rb") as f:
        return pickle.load(f)


def predicted_time(test_time):
    # same rounding as the released code: next multiple of 5 s plus the window, capped at the test duration
    t = math.ceil(test_time / PREDICTION_STEP_SEC) * PREDICTION_STEP_SEC + PREDICTION_WINDOW_SEC
    return min(t, config.TEST_DURATION)


def predict(model, at_sec, **regressors):
    future = pd.DataFrame({"ds": [pd.to_datetime(at_sec, unit="s")], **{k: [v] for k, v in regressors.items()}})
    yhat = model.predict(future)["yhat"].iloc[0]
    return math.ceil(yhat)


#************************************************************************** Microservice Manager (Algorithm 3) **************************************************************************

_demand_models = {}


def _init_worker():
    for service in MANAGERS.values():
        _demand_models[service] = load_model(service, "demand")


def microservice_manager(args):
    module_name, test_time = args
    manager = __import__(module_name)

    # Monitor_M: unchanged Smart HPA Monitor
    name, spec_replicas, current_replicas, current_cpu, target_cpu, cpu_request, max_replica, min_replica = manager.Monitor()

    # reactive DR for the Knowledge Base (column D), same formula as Smart HPA's Analyse
    cpu_percentage = (current_cpu / cpu_request) * 100
    reactive_dr = math.ceil(int(current_replicas) * (int(cpu_percentage) / int(target_cpu)))

    # Analyze_M, line 1: PDR = PROPHET(RM, RMT, CR); RMT is constant in a scenario so it is not a model input
    p_time = predicted_time(test_time)
    if test_time > PREDICTION_START_SEC:
        pdr = predict(_demand_models[name], p_time, cpu_pct=cpu_percentage, current_replicas=current_replicas)
    else:
        pdr = current_replicas

    # Analyze_M, lines 2-8
    if pdr > current_replicas:
        sd = "scale up"
    elif pdr < current_replicas and pdr >= min_replica:
        sd = "scale down"
    else:
        sd = "no scale"

    return {"name": name, "cpu_pct": cpu_percentage, "cr": current_replicas, "reactive_dr": reactive_dr,
            "pdr": pdr, "sd": sd, "maxr": max_replica, "minr": min_replica, "req": cpu_request, "p_time": p_time}


#*********************************************************************** Application Resource Manager (Algorithm 4) **********************************************************************

def _added_removed(ms):
    added = sum((m["pmaxr"] - m["maxr"]) * m["req"] for m in ms if m["pmaxr"] > m["maxr"])
    removed = sum((m["maxr"] - m["pmaxr"]) * m["req"] for m in ms if m["pmaxr"] < m["maxr"])
    return added, removed


def balance(ms):
    """Plan_A, lines 2-14: make the resource added by PROPHET equal (at most) the resource it removed, so the
    application's total quota stays the SLA total. Works in whole replicas, so with different requests per service an
    exact match may not exist; the result then stays slightly BELOW the total, never above."""
    added, removed = _added_removed(ms)

    if added > removed:
        # line 9: extract from over-provisioned services, most over-provisioned first
        for m in sorted(ms, key=lambda m: (m["pmaxr"] - m["pdr"]) * m["req"], reverse=True):
            need = added - removed
            spare = m["pmaxr"] - max(m["pdr"], m["minr"])
            if need <= 0 or spare <= 0:
                continue
            m["pmaxr"] -= min(spare, math.ceil(need / m["req"]))
            added, removed = _added_removed(ms)
        # still unbalanced: PROPHET added more than the others can give; trim the additions of the least
        # under-provisioned services first, keeping the most under-provisioned ones (paper: "prioritizing highly
        # under-provisioned microservices")
        for m in sorted(ms, key=lambda m: (m["pdr"] - m["maxr"]) * m["req"]):
            need = added - removed
            extra = m["pmaxr"] - m["maxr"]
            if need <= 0 or extra <= 0:
                continue
            m["pmaxr"] -= min(extra, math.ceil(need / m["req"]))
            added, removed = _added_removed(ms)

    elif removed > added:
        # line 11: add to under-provisioned services, most under-provisioned first
        for m in sorted(ms, key=lambda m: (m["pdr"] - m["pmaxr"]) * m["req"], reverse=True):
            free = removed - added
            short = m["pdr"] - m["pmaxr"]
            if short <= 0:
                continue
            m["pmaxr"] += min(short, free // m["req"])
            added, removed = _added_removed(ms)
        # quota nobody needs goes back to the services it was taken from, so the application total is kept
        for m in ms:
            free = removed - added
            if m["pmaxr"] < m["maxr"]:
                m["pmaxr"] += min(m["maxr"] - m["pmaxr"], free // m["req"])
                added, removed = _added_removed(ms)
    return ms


def application_resource_manager(ms, capacity_models):
    # Analyze_A, line 1: PmaxR = PROPHET(PDR, maxR, CR); maxR is the constant SLA value so it is not a model input
    for m in ms:
        m["pmaxr"] = max(m["minr"], predict(capacity_models[m["name"]], m["p_time"], desired=m["pdr"], current_replicas=m["cr"]))

    balance(ms)

    # Execute_A, lines 15-21
    for m in ms:
        if m["pmaxr"] >= m["pdr"]:
            m["res_sd"], m["res_dr"] = m["sd"], m["pdr"]
        elif m["cr"] < m["pmaxr"] < m["pdr"]:
            m["res_sd"], m["res_dr"] = "scale up", m["pmaxr"]
        else:
            m["res_sd"], m["res_dr"] = "no scale", m["cr"]
    return ms


#****************************************************************************** Execute and Knowledge Base ******************************************************************************

def execute(ms):
    # kubectl scale for every service that changes, in parallel (Smart HPA used one process per scale command)
    procs = [subprocess.Popen(["kubectl", "scale", "deployment", m["name"], f"--replicas={m['res_dr']}"])
             for m in ms if m["res_sd"] != "no scale"]
    for p in procs:
        p.wait()


KB_HEADER = ["Test Time (sec)", "CPU Usage Percentage", "Current Replicas", "Desired Replicas", "Max. Replicas",
             "Scaling Action", "Predicted Desired Replicas", "Executed Replicas", "Predicted Time (sec)", "ARM Active"]


def store(ms, test_time, row, arm_active):
    for m in ms:
        path = f"./Knowledge Base/{m['name']}.xlsx"
        workbook = load_workbook(path)
        sheet = workbook.active
        for col, title in enumerate(KB_HEADER, start=1):
            sheet.cell(row=1, column=col, value=title)
        values = [test_time, m["cpu_pct"], m["cr"], m["reactive_dr"], m["pmaxr"], m["res_sd"],
                  m["pdr"], m["res_dr"], m["p_time"], int(arm_active)]
        for col, value in enumerate(values, start=1):
            sheet.cell(row=row, column=col, value=value)
        workbook.save(path)


#************************************************************************************* Main loop *************************************************************************************

if __name__ == "__main__":
    capacity_models = {service: load_model(service, "capacity") for service in MANAGERS.values()}

    # one pool for the whole run: each worker loads PROPHET and the 11 demand models once
    with multiprocessing.Pool(processes=len(MANAGERS), initializer=_init_worker) as pool:
        start_time = time.time()
        row = 2
        while (time.time() - start_time) < config.TEST_DURATION:
            test_time = time.time() - start_time

            ms = pool.map(microservice_manager, [(module, test_time) for module in MANAGERS])

            # Plan_M, line 9: any predicted demand above capacity activates the ARM (once for all services)
            arm_active = any(m["pdr"] > m["maxr"] for m in ms)
            if arm_active:
                application_resource_manager(ms, capacity_models)
            else:
                for m in ms:
                    m["pmaxr"], m["res_sd"], m["res_dr"] = m["maxr"], m["sd"], m["pdr"]
            for m in ms:
                if m["res_sd"] == "no scale":
                    m["res_dr"] = m["cr"]               # nothing is executed, so the replica count stays as it is

            execute(ms)
            store(ms, test_time, row, arm_active)
            row += 1

            print(f"{test_time:6.1f}s ARM={int(arm_active)} " +
                  " ".join(f"{m['name']}:{m['cr']}->{m['res_dr']}/{m['pmaxr']}(pdr {m['pdr']})" for m in ms), flush=True)
