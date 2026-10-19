# Scaling stability of one scenario (plan.md section 7, improvement C2): does the autoscaler move replicas and quota
# back and forth when the load jumps up and down?
# Reads results/<scenario>/<autoscaler>/run<N>/knowledge_base/*.xlsx (one row per loop iteration, written by Smart HPA
# and by k8shpa_observer.py alike) and writes results/<scenario>/comparison/flapping.csv and flapping.md.
#
# Per run, summed over the 11 microservices:
#   replica changes    iterations where the current replica count (CR) differs from the previous iteration
#   scale reversals    a scale-up followed by a scale-down of the same service (or the reverse): flapping proper
#   maxR changes       iterations where maxR differs from the previous iteration, i.e. quota moved between services
#                      (always 0 for K8s HPA, whose maxR is fixed)
#   max replicas swing largest (highest - lowest) CR of a service after its first scale-up
#
# Usage: python make_flapping.py <results/<scenario> dir> [--include-trials] [--c3 | --c3-control]

import csv
import os
import sys

from make_comparison import (AUTOSCALERS, COMPARISONS, NAME, consistency_lines, means_lines, output_folder, run_dirs,
                             scenario_label)
from make_run_report import SERVICES
from openpyxl import load_workbook

# (key, label, True if higher is better)
METRICS = [
    ("replica_changes", "Replica changes (count)", False),
    ("scale_reversals", "Scale reversals (count)", False),
    ("maxr_changes", "maxR changes / quota moves (count)", False),
    ("max_swing", "Largest replica swing of one service", False),
]


def kb_series(run_dir, service):
    """(CR, maxR) per loop iteration."""
    rows = list(load_workbook(os.path.join(run_dir, "knowledge_base", f"{service}.xlsx")).active.iter_rows(values_only=True))[1:]
    rows = [r for r in rows if r[0] is not None and r[4] is not None]   # last iteration may miss columns E/F
    return [r[2] for r in rows], [r[4] for r in rows]


def stability(run_dir):
    m = {k: 0 for k, _, _ in METRICS}
    for s in SERVICES:
        cr, maxr = kb_series(run_dir, s)
        steps = [b - a for a, b in zip(cr, cr[1:]) if b != a]
        m["replica_changes"] += len(steps)
        m["scale_reversals"] += sum(1 for a, b in zip(steps, steps[1:]) if (a > 0) != (b > 0))
        m["maxr_changes"] += sum(1 for a, b in zip(maxr, maxr[1:]) if b != a)
        first_up = next((i for i in range(1, len(cr)) if cr[i] > cr[i - 1]), None)
        if first_up is not None:
            m["max_swing"] = max(m["max_swing"], max(cr[first_up:]) - min(cr[first_up:]))
    return m


def main():
    scenario_dir = os.path.normpath(sys.argv[1])
    scenario = scenario_label(scenario_dir)
    out_dir = os.path.join(scenario_dir, output_folder(sys.argv))
    runs = {(a, n): stability(d) for a, n, d in run_dirs(scenario_dir, "--include-trials" in sys.argv)}
    if not runs:
        print(f"No runs under {scenario_dir}")
        return
    present = [a for a, _ in AUTOSCALERS if any(aa == a for aa, _ in runs)]
    comparisons = [(a, b) for a, b in COMPARISONS if a in present and b in present]

    os.makedirs(out_dir, exist_ok=True)
    keys = [k for k, _, _ in METRICS]
    with open(os.path.join(out_dir, "flapping.csv"), "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(["scenario", "autoscaler", "run"] + keys)
        for (a, n), m in sorted(runs.items(), key=lambda x: (x[0][1], x[0][0])):
            w.writerow([scenario, a, n] + [m[k] for k in keys])

    lines = [f"# Scaling stability: {scenario}", "",
             f"From the Knowledge Base of {len(runs)} runs, summed over the 11 microservices. A scale reversal is a "
             f"scale-up followed by a scale-down of the same service, or the reverse. maxR changes count quota moves "
             f"(0 for K8s HPA, whose maxR is fixed)."]
    lines += consistency_lines(runs, comparisons, METRICS) + means_lines(runs, present, METRICS)
    with open(os.path.join(out_dir, "flapping.md"), "w", encoding="utf-8") as f:
        f.write("\n".join(lines) + "\n")
    print(f"Scaling stability written to {out_dir}")


if __name__ == "__main__":
    main()
