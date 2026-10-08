#*********************************************************** Knowledge Base logger for the Kubernetes HPA baseline ***********************************************************
#
# The Smart HPA package only records data for Smart HPA. For the Kubernetes HPA baseline we run the SAME Microservice
# Manager code (Monitor + Analyse + Knowledge Base writes), but never scale anything and never call the Adaptive Resource
# Manager: the built-in Kubernetes HPA does the scaling. This mirrors the resource-rich branch of
# Microservice_Capacity_Analyzer.py, so both autoscalers are measured by identical code and produce identical xlsx files.
#
# Column D (Desired Replicas) is therefore ceil(CR x CPU% / threshold) and is NOT capped by maxR, which is what the
# Results Analysis script needs to compute CPU Underprovision. Column E is the fixed SLA maxR, because Kubernetes HPA
# never changes it.

import time
import multiprocessing
from functools import partial
from openpyxl import load_workbook
from frontend import *
from adservice import *
from checkoutservice import *
from currencyservice import *
from emailservice import *
from paymentservice import *
from productcatalogservice import *
from cartservice import *
from recommendationservice import *
from shippingservice import *
from rediscart import *
import config


def run_function(func):
    return func()


desired_time = config.TEST_DURATION
start_time = time.time()

if __name__ == '__main__':

    row_number = 2                    # same row bookkeeping as Microservice_Capacity_Analyzer.py

    while (time.time() - start_time) < desired_time:

        Test_Time = time.time() - start_time

        functions = [partial(frontend, Test_Time), partial(adservice, Test_Time), partial(cartservice, Test_Time), partial(currencyservice, Test_Time), partial(checkoutservice, Test_Time), partial(emailservice, Test_Time), partial(paymentservice, Test_Time), partial(shippingservice, Test_Time), partial(productcatalogservice, Test_Time), partial(recommendationservice, Test_Time), partial(rediscart, Test_Time)]
        with multiprocessing.Pool(processes=len(functions)) as pool:
            microservices_data = pool.map(run_function, functions)

        for i in range(len(microservices_data)):
            filename = microservices_data[i][0]
            workbook = load_workbook(f'./Knowledge Base/{filename}.xlsx')
            sheet = workbook.active
            sheet.cell(row=row_number, column=5, value=microservices_data[i][5])            # SLA maxR (Kubernetes HPA never changes it)
            sheet.cell(row=row_number, column=6, value=microservices_data[i][1])            # scaling decision Smart HPA's manager would make (not executed)
            workbook.save(f'./Knowledge Base/{filename}.xlsx')

        row_number = row_number + 1

        print(f"[{Test_Time:7.1f}s] observed", flush=True)
