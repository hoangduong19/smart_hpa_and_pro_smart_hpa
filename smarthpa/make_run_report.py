# Builds the per-run report from one results folder (results/<scenario>/<autoscaler>/run<N>/):
#   metrics.csv      Table 1 metrics per microservice + application level
#   summary.md       the same numbers plus run settings and Locust totals, readable on its own
#   fig_demand_capacity.png   CPU demand vs capacity vs supply over time, one panel per microservice (paper Fig. 7a/b)
#   fig_cpu_utilization.png   CPU utilization vs threshold over time, one panel per microservice (paper Fig. 7c/d)
#   fig_load.png              Locust users, requests/s, p95 latency and failures/s over time
#
# Metric formulas follow "Results Analysis Script/Results_Analysis.py" of the Smart HPA package:
#   Supply = mean(CR x request), Overprovision = mean(max(maxR - DR, 0) x request),
#   Underprovision = mean(max(DR - maxR, 0) x request), Overutilization = mean(CPU% if CPU% > threshold else 0).
# The time metrics are re-implemented: the original loops index into empty lists and cannot run. Here a time metric is
# the sum of loop intervals (t[i+1] - t[i]) where the condition holds, and the application-level times follow Table 1
# of the paper (over-utilization / under-provision: at least one microservice; over-provision: no microservice
# under-provisioned).
#
# Usage: python make_run_report.py <run_dir>

import csv
import json
import os
import statistics
import sys

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from openpyxl import load_workbook

SERVICES = ["frontend", "currencyservice", "cartservice", "adservice", "checkoutservice", "productcatalogservice",
            "recommendationservice", "paymentservice", "shippingservice", "emailservice", "redis-cart"]
CPU_REQUEST = {"adservice": 200, "cartservice": 200, "redis-cart": 70}          # all other services: 100m

# Reference palette (dataviz skill, light mode): three series per panel = slots 1-3, validated all-pairs.
SURFACE, INK, INK_2, GRID = "#fcfcfb", "#0b0b0b", "#52514e", "#e6e5e0"
DEMAND, CAPACITY, SUPPLY = "#2a78d6", "#eb6834", "#1baf7a"


def load_kb(run_dir, service):
    rows = list(load_workbook(os.path.join(run_dir, "knowledge_base", f"{service}.xlsx")).active.iter_rows(values_only=True))[1:]
    rows = [r for r in rows if r[0] is not None and r[4] is not None]   # last iteration may miss columns E/F
    req = CPU_REQUEST.get(service, 100)
    return {
        "t": [r[0] for r in rows],
        "cpu_pct": [r[1] for r in rows],
        "demand": [r[3] * req for r in rows],
        "capacity": [r[4] * req for r in rows],
        "supply": [r[2] * req for r in rows],
    }


def duration(t, flags):
    return sum(t[i + 1] - t[i] for i in range(len(t) - 1) if flags[i])


def compute_metrics(kb, threshold):
    per_service = {}
    for s, d in kb.items():
        over_res = [max(c - dm, 0) for c, dm in zip(d["capacity"], d["demand"])]
        under_res = [max(dm - c, 0) for c, dm in zip(d["capacity"], d["demand"])]
        overutil = [p if p > threshold else 0 for p in d["cpu_pct"]]
        per_service[s] = {
            "supply_cpu_m": statistics.mean(d["supply"]),
            "cpu_overutilization_pct": statistics.mean(overutil),
            "overutilization_time_min": duration(d["t"], [v > 0 for v in overutil]) / 60,
            "cpu_overprovision_m": statistics.mean(over_res),
            "overprovision_time_min": duration(d["t"], [v > 0 for v in over_res]) / 60,
            "cpu_underprovision_m": statistics.mean(under_res),
            "underprovision_time_min": duration(d["t"], [v > 0 for v in under_res]) / 60,
        }

    # Application level: sums for resource metrics, Table 1 definitions for times (rows are aligned: every manager
    # writes one row per loop iteration).
    n = min(len(d["t"]) for d in kb.values())
    t = next(iter(kb.values()))["t"][:n]
    any_overutil = [any(kb[s]["cpu_pct"][i] > threshold for s in kb) for i in range(n)]
    any_under = [any(kb[s]["demand"][i] > kb[s]["capacity"][i] for s in kb) for i in range(n)]
    app = {k: sum(m[k] for m in per_service.values()) for k in
           ["supply_cpu_m", "cpu_overutilization_pct", "cpu_overprovision_m", "cpu_underprovision_m"]}
    app["overutilization_time_min"] = duration(t, any_overutil) / 60
    app["underprovision_time_min"] = duration(t, any_under) / 60
    app["overprovision_time_min"] = duration(t, [not u for u in any_under]) / 60
    return per_service, app


def read_locust(run_dir):
    hist = os.path.join(run_dir, "locust_stats_history.csv")
    rows = [r for r in csv.DictReader(open(hist, newline="")) if r["Name"] == "Aggregated"]
    t0 = int(rows[0]["Timestamp"])
    num = lambda v: float(v) if v not in ("N/A", "") else None
    series = {
        "t": [(int(r["Timestamp"]) - t0) / 60 for r in rows],
        "users": [num(r["User Count"]) for r in rows],
        "rps": [num(r["Requests/s"]) for r in rows],
        "p95": [num(r["95%"]) for r in rows],
        "fails": [num(r["Failures/s"]) for r in rows],
    }
    total = [r for r in csv.DictReader(open(os.path.join(run_dir, "locust_stats.csv"), newline="")) if r["Name"] == "Aggregated"][0]
    return series, total


def style_axis(ax):
    ax.set_facecolor(SURFACE)
    for side in ("top", "right"):
        ax.spines[side].set_visible(False)
    for side in ("left", "bottom"):
        ax.spines[side].set_color(GRID)
    ax.tick_params(colors=INK_2, labelsize=8, length=0)
    ax.grid(axis="y", color=GRID, linewidth=0.8)
    ax.set_axisbelow(True)


def panel_grid(title):
    fig, axes = plt.subplots(3, 4, figsize=(16, 9.5), sharex=True, facecolor=SURFACE)
    fig.suptitle(title, x=0.01, ha="left", fontsize=14, color=INK, fontweight="bold")
    return fig, axes.flatten()


def fig_demand_capacity(kb, run_label, out):
    fig, axes = panel_grid(f"CPU demand vs capacity vs supply (milliCPU) - {run_label}")
    for ax, s in zip(axes, SERVICES):
        d = kb[s]
        t = [x / 60 for x in d["t"]]
        style_axis(ax)
        ax.fill_between(t, d["supply"], step="post", color=SUPPLY, alpha=0.18, linewidth=0)
        ax.step(t, d["supply"], where="post", color=SUPPLY, linewidth=1.5, label="Supply (current replicas)")
        ax.step(t, d["demand"], where="post", color=DEMAND, linewidth=2, label="Demand (desired replicas)")
        ax.step(t, d["capacity"], where="post", color=CAPACITY, linewidth=2, linestyle=(0, (4, 2)), label="Capacity (maxR)")
        ax.set_title(s, fontsize=10, color=INK, loc="left")
        ax.set_ylim(bottom=0)
    legend_ax = axes[-1]
    legend_ax.axis("off")
    handles, labels = axes[0].get_legend_handles_labels()
    legend_ax.legend(handles, labels, loc="center left", frameon=False, fontsize=10, labelcolor=INK)
    for ax in axes[8:11]:
        ax.set_xlabel("Test time (min)", fontsize=9, color=INK_2)
    fig.tight_layout(rect=(0, 0, 1, 0.96))
    fig.savefig(out, dpi=130, facecolor=SURFACE)
    plt.close(fig)


def fig_cpu_utilization(kb, threshold, run_label, out):
    fig, axes = panel_grid(f"CPU utilization (% of request) vs {threshold}% threshold - {run_label}")
    for ax, s in zip(axes, SERVICES):
        d = kb[s]
        t = [x / 60 for x in d["t"]]
        style_axis(ax)
        ax.plot(t, d["cpu_pct"], color=DEMAND, linewidth=2, label="CPU utilization")
        ax.axhline(threshold, color=INK_2, linewidth=1.2, linestyle=(0, (4, 2)), label=f"Threshold {threshold}%")
        ax.set_title(s, fontsize=10, color=INK, loc="left")
        ax.set_ylim(bottom=0)
    legend_ax = axes[-1]
    legend_ax.axis("off")
    handles, labels = axes[0].get_legend_handles_labels()
    legend_ax.legend(handles, labels, loc="center left", frameon=False, fontsize=10, labelcolor=INK)
    for ax in axes[8:11]:
        ax.set_xlabel("Test time (min)", fontsize=9, color=INK_2)
    fig.tight_layout(rect=(0, 0, 1, 0.96))
    fig.savefig(out, dpi=130, facecolor=SURFACE)
    plt.close(fig)


def fig_load(loc, run_label, out):
    panels = [("users", "Concurrent users"), ("rps", "Requests / s"), ("p95", "p95 response time (ms)"), ("fails", "Failures / s")]
    fig, axes = plt.subplots(4, 1, figsize=(10, 9), sharex=True, facecolor=SURFACE)
    fig.suptitle(f"Load test (Locust) - {run_label}", x=0.01, ha="left", fontsize=14, color=INK, fontweight="bold")
    for ax, (key, title) in zip(axes, panels):
        style_axis(ax)
        ax.plot(loc["t"], [v if v is not None else float("nan") for v in loc[key]], color=DEMAND, linewidth=2)
        ax.set_title(title, fontsize=10, color=INK, loc="left")
        ax.set_ylim(bottom=0)
    axes[-1].set_xlabel("Test time (min)", fontsize=9, color=INK_2)
    fig.tight_layout(rect=(0, 0, 1, 0.96))
    fig.savefig(out, dpi=130, facecolor=SURFACE)
    plt.close(fig)


def fmt(v, digits=2):
    return f"{v:,.{digits}f}"


def main():
    run_dir = sys.argv[1]
    info = json.load(open(os.path.join(run_dir, "run_info.json"), encoding="utf-8-sig"))
    threshold = int(info["cpu_threshold"])
    run_label = f"{info['autoscaler']} | {info['scenario']} | run {info['run']}"

    kb = {s: load_kb(run_dir, s) for s in SERVICES}
    per_service, app = compute_metrics(kb, threshold)
    loc, total = read_locust(run_dir)

    keys = ["supply_cpu_m", "cpu_overutilization_pct", "overutilization_time_min", "cpu_overprovision_m",
            "overprovision_time_min", "cpu_underprovision_m", "underprovision_time_min"]
    with open(os.path.join(run_dir, "metrics.csv"), "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(["autoscaler", "scenario", "run", "microservice"] + keys)
        for s in SERVICES:
            w.writerow([info["autoscaler"], info["scenario"], info["run"], s] + [round(per_service[s][k], 3) for k in keys])
        w.writerow([info["autoscaler"], info["scenario"], info["run"], "APPLICATION"] + [round(app[k], 3) for k in keys])

    unsched = 0
    rec = os.path.join(run_dir, "cluster_metrics.csv")
    if os.path.exists(rec):
        unsched = max((int(r.get("unschedulable_pods_total") or 0) for r in csv.DictReader(open(rec, newline=""))), default=0)

    labels = {
        "supply_cpu_m": "Supply CPU (m) - higher is better",
        "cpu_overutilization_pct": "CPU Overutilization (%) - lower",
        "overutilization_time_min": "Overutilization Time (min) - lower",
        "cpu_overprovision_m": "CPU Overprovision (m) - lower",
        "overprovision_time_min": "Overprovision Time (min) - higher",
        "cpu_underprovision_m": "CPU Underprovision (m) - lower",
        "underprovision_time_min": "Underprovision Time (min) - lower",
    }
    lines = [
        f"# Run report: {run_label}",
        "",
        f"- Started {info.get('started_at')} - finished {info.get('finished_at')}",
        f"- Scenario {info['scenario']}: maxR = {info['max_replicas']}, CPU threshold = {threshold}%",
        f"- Load: {info['users']} users, spawn rate {info['spawn_rate']}/s, {info['duration_sec']} s",
        f"- Knowledge Base rows: {min(len(d['t']) for d in kb.values())} loop iterations",
        f"- Max unschedulable pods during run: {unsched}",
        "",
        "## Application level (Table 1 metrics)",
        "",
        "| Metric | Value |",
        "|---|---|",
    ] + [f"| {labels[k]} | {fmt(app[k])} |" for k in keys] + [
        "",
        "## Load test totals (Locust)",
        "",
        "| Requests | Failures | Median (ms) | p95 (ms) | p99 (ms) | Requests/s |",
        "|---|---|---|---|---|---|",
        f"| {total['Request Count']} | {total['Failure Count']} | {total['Median Response Time']} | {total['95%']} | {total['99%']} | {fmt(float(total['Requests/s']))} |",
        "",
        "## Per microservice",
        "",
        "| Microservice | Supply (m) | Overutil (%) | Overutil time (min) | Overprov (m) | Overprov time (min) | Underprov (m) | Underprov time (min) |",
        "|---|---|---|---|---|---|---|---|",
    ] + [f"| {s} | " + " | ".join(fmt(per_service[s][k]) for k in keys) + " |" for s in SERVICES] + [
        "",
        "Figures: `fig_demand_capacity.png`, `fig_cpu_utilization.png`, `fig_load.png`.",
    ]
    with open(os.path.join(run_dir, "summary.md"), "w", encoding="utf-8") as f:
        f.write("\n".join(lines) + "\n")

    fig_demand_capacity(kb, run_label, os.path.join(run_dir, "fig_demand_capacity.png"))
    fig_cpu_utilization(kb, threshold, run_label, os.path.join(run_dir, "fig_cpu_utilization.png"))
    fig_load(loc, run_label, os.path.join(run_dir, "fig_load.png"))
    print(f"Report written to {run_dir}")


if __name__ == "__main__":
    main()
